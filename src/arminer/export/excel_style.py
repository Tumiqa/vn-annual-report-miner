# -*- coding: utf-8 -*-
"""
arminer.export.excel_style
===========================
Shared premium styling and branding utility for ALL Excel exports.
Cung cấp bộ nhận diện thương hiệu và styling chuyên nghiệp nhất quán cho mọi file .xlsx xuất ra từ arminer:
- Tự động tạo Trang_Bia (Cover Sheet) với Logo, thông tin Trương Minh Quân, mục lục, trích dẫn khoa học.
- Header hiện đại, zebra striping mềm mại, căn chỉnh chuẩn kinh tế lượng, định dạng số khoa học.
- Khóa cột cố định (Freeze Panes cho Ticker, Year), bộ lọc tự động (AutoFilter), tự động giãn độ rộng cột.
- Cài đặt thông tin bản quyền (Properties), Header & Footer khi in ấn / xuất PDF.

Usage:
    from arminer.export.excel_style import style_excel_file
    style_excel_file("path/to/file.xlsx")
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
from openpyxl.utils import get_column_letter
import pandas as pd
from loguru import logger


# =====================================================================
# PREMIUM COLOR PALETTE
# =====================================================================

NAVY = "1B3A5C"
DARK_SLATE = "0F172A"
TEAL = "0D7377"
EMERALD = "059669"
GOLD = "D4A843"
PURPLE = "7C3AED"
GREEN = "2E7D32"
DARK_BLUE = "205375"
WHITE = "FFFFFF"
LIGHT_GRAY = "F8FAFC"
HEADER_GRAY = "F1F5F9"
BORDER_COLOR = "CBD5E1"
BORDER_LIGHT = "E2E8F0"

# Pre-built fills
FILL_NAVY = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
FILL_TEAL = PatternFill(start_color=TEAL, end_color=TEAL, fill_type="solid")
FILL_PURPLE = PatternFill(start_color=PURPLE, end_color=PURPLE, fill_type="solid")
FILL_GREEN = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")
FILL_DARK_BLUE = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
FILL_WHITE = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
FILL_ZEBRA = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
FILL_SECTION_HDR = PatternFill(start_color=HEADER_GRAY, end_color=HEADER_GRAY, fill_type="solid")

# Pre-built fonts
FONT_HEADER = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
FONT_BODY = Font(name="Segoe UI", size=10, color="1E293B")
FONT_BODY_BOLD = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
FONT_SMALL = Font(name="Segoe UI", size=9, color="64748B")
FONT_CODE = Font(name="Consolas", size=9.5, color="0F172A")

# Pre-built borders
THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER_LIGHT),
    right=Side(style="thin", color=BORDER_LIGHT),
    top=Side(style="thin", color=BORDER_LIGHT),
    bottom=Side(style="thin", color=BORDER_LIGHT),
)

HEADER_BORDER = Border(
    left=Side(style="thin", color="475569"),
    right=Side(style="thin", color="475569"),
    top=Side(style="thin", color="475569"),
    bottom=Side(style="medium", color="0F172A"),
)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

# Color palette per sheet name for auto-selection
_SHEET_COLORS = {
    "Trang_Bia": GOLD,
    "Cover_Sheet": GOLD,
    "Firm_Year_Panel": NAVY,
    "Panel_Data": NAVY,
    "Merged_Panel": DARK_BLUE,
    "Articles_Detail": TEAL,
    "Raw_Keywords": PURPLE,
    "Context": EMERALD,
    "Codebook": DARK_BLUE,
    "Financial_Data": NAVY,
    "Bao_Cao_Tai_Chinh": NAVY,
    "Ty_So_Tai_Chinh": TEAL,
    "Company_Info": EMERALD,
}

# Standard sheet description lookup for Cover Sheet
_SHEET_DESCRIPTIONS = {
    "Firm_Year_Panel": "Bảng dữ liệu bảng cấp độ Doanh nghiệp - Năm (Biến tần suất từ khóa, Log, Mention, Mật độ)",
    "Panel_Data": "Bảng dữ liệu bảng cấp độ Doanh nghiệp - Năm phục vụ mô hình hồi quy kinh tế lượng",
    "Articles_Detail": "Danh sách chi tiết từng bài báo / văn bản đã thu thập (Tiêu đề, nguồn tin, ngày, số từ, URL)",
    "Raw_Keywords": "Thống kê chi tiết tần suất xuất hiện của từng từ khóa đơn lẻ và phân loại nhóm chủ đề",
    "Context": "Ngữ cảnh câu chứa từ khóa — Trích xuất thông minh theo ranh giới câu (±1 câu lân cận) cho từng lần xuất hiện",
    "Codebook": "Từ điển giải thích chi tiết ý nghĩa, thang đo, công thức tính toán và tài liệu tham khảo",
    "Merged_Panel": "Dữ liệu bảng ghép nối hoàn chỉnh giữa biến văn bản (Mining) và biến tài chính (BCTC)",
    "Bao_Cao_Tai_Chinh": "Toàn bộ 702 chỉ tiêu kế toán chi tiết theo 13 nhóm chuẩn mực kế toán Việt Nam",
    "Ty_So_Tai_Chinh": "Hệ thống 116 chỉ số tài chính chuyên sâu chuẩn học thuật (Sinh lời, Cấu trúc vốn, Thanh khoản, Dòng tiền, CAMELS, Altman Z-score)",
    "Company_Info": "Danh sách toàn bộ doanh nghiệp niêm yết trên HOSE, HNX, UPCoM — Phân ngành ICB 4 cấp chuẩn FiinPro",
}


def _pick_header_fill(sheet_name: str) -> PatternFill:
    """Auto-pick a header fill color based on sheet name."""
    color = _SHEET_COLORS.get(sheet_name, NAVY)
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _auto_col_width(ws, min_width: int = 11, max_width: int = 60) -> None:
    """Auto-size column widths based on content with generous padding."""
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in col_cells[:120]:  # sample first 120 rows
            if cell.value is not None:
                val_str = str(cell.value)
                cell_len = len(val_str)
                if cell_len > max_len:
                    max_len = cell_len
        width = min(max(max_len + 4, min_width), max_width)
        ws.column_dimensions[col_letter].width = width


def _guess_number_format(col_name: str) -> Optional[str]:
    """Guess appropriate number format from column name."""
    cl = col_name.lower().strip()
    if cl in ("year", "nam", "published_year"):
        return "0"
    if cl in ("mention", "dummy", "substantive"):
        return "0"
    if any(k in cl for k in ("pct", "ratio", "rate", "roa", "roe", "ros", "margin", "ty_le", "bien")):
        return "0.00%"
    if any(k in cl for k in ("log", "ln_", "log_frequency", "log_freq")):
        return "0.0000"
    if any(k in cl for k in ("density", "mat_do", "coverage")):
        return "0.0000"
    if cl in ("n", "obs", "n_obs", "observations", "so_luong", "total_words", "word_count", "hits", "total_articles", "articles_with_hits", "total_mentions", "unique_keywords"):
        return "#,##0"
    if any(k in cl for k in ("_freq", "count", "so_luong")) or cl.endswith("frequency") or cl == "frequency":
        return "#,##0"
    if cl in ("mean", "std dev", "std", "median", "min", "max"):
        return "#,##0.0000"
    return None


def add_cover_sheet(wb: openpyxl.Workbook, custom_title: Optional[str] = None) -> openpyxl.worksheet.worksheet.Worksheet:
    """
    Tạo hoặc bổ sung Trang_Bia (Cover Sheet) đại diện thương hiệu arminer cho toàn bộ file Excel.
    Bao gồm:
    - Logo arminer và thông tin bản quyền Trương Minh Quân
    - Thời gian xuất dữ liệu và tổng quan kỹ thuật
    - Mục lục giải thích chi tiết mục đích từng Sheet
    """
    # If Trang_Bia or Cover_Sheet already exists, return existing
    for sname in ("Trang_Bia", "Cover_Sheet", "Cover"):
        if sname in wb.sheetnames:
            return wb[sname]

    ws = wb.create_sheet(title="Trang_Bia")
    # Move to index 0
    wb._sheets.insert(0, wb._sheets.pop(wb._sheets.index(ws)))
    wb.active = ws

    ws.sheet_properties.tabColor = GOLD
    try:
        ws.views.sheetView[0].showGridLines = True
    except Exception:
        pass

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 58
    ws.column_dimensions["D"].width = 28

    # Set row heights
    ws.row_dimensions[1].height = 14
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 28
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 18
    ws.row_dimensions[6].height = 12

    # Try embedding PNG logo image
    logo_path = Path(__file__).resolve().parent.parent / "ui" / "static" / "arminer_logo.png"
    has_image = False
    if logo_path.exists():
        try:
            img = openpyxl.drawing.image.Image(str(logo_path))
            img.width = 68
            img.height = 68
            ws.add_image(img, "B3")
            has_image = True
        except Exception as e:
            logger.debug(f"Cover sheet logo insertion skipped: {e}")

    # Header Titles
    title_text = custom_title or "ARMINER WEB STUDIO"
    if has_image:
        ws.cell(row=3, column=3, value=title_text).font = Font(name="Segoe UI", size=18, bold=True, color=DARK_SLATE)
        ws.cell(row=4, column=3, value="HỆ THỐNG KHAI PHÁ DỮ LIỆU DOANH NGHIỆP & NGHIÊN CỨU ĐỊNH LƯỢNG").font = Font(
            name="Segoe UI", size=11, bold=True, color=TEAL
        )
        ws.cell(row=5, column=3, value="Vietnam Corporate Text Mining & Econometric Intelligence Platform").font = Font(
            name="Segoe UI", size=9.5, italic=True, color="64748B"
        )
    else:
        ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=4)
        ws.cell(row=3, column=2, value=title_text).font = Font(name="Segoe UI", size=18, bold=True, color=DARK_SLATE)
        ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=4)
        ws.cell(row=4, column=2, value="HỆ THỐNG KHAI PHÁ DỮ LIỆU DOANH NGHIỆP & NGHIÊN CỨU ĐỊNH LƯỢNG").font = Font(
            name="Segoe UI", size=11, bold=True, color=TEAL
        )
        ws.merge_cells(start_row=5, start_column=2, end_row=5, end_column=4)
        ws.cell(row=5, column=2, value="Vietnam Corporate Text Mining & Econometric Intelligence Platform").font = Font(
            name="Segoe UI", size=9.5, italic=True, color="64748B"
        )

    # Gold Accent Divider line
    for col in (2, 3, 4):
        ws.cell(row=6, column=col).border = Border(bottom=Side(style="medium", color=GOLD))

    curr_row = 8

    # --- Section I: Thông tin bộ dữ liệu & Bản quyền ---
    ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=4)
    sec1 = ws.cell(row=curr_row, column=2, value="I. THÔNG TIN BỘ DỮ LIỆU & BẢN QUYỀN CÔNG CỤ")
    sec1.font = Font(name="Segoe UI", size=11, bold=True, color=NAVY)
    sec1.fill = FILL_SECTION_HDR
    sec1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in (2, 3, 4):
        ws.cell(row=curr_row, column=col).border = THIN_BORDER
    ws.row_dimensions[curr_row].height = 24
    curr_row += 1

    info_items = [
        ("🏛️ Đơn vị phát triển:", "Trương Minh Quân - Đại học Kinh tế Đà Nẵng"),
        ("⚙️ Công cụ trích xuất:", "arminer Studio v2.5 (Corporate Text Mining & Econometric Automation)"),
        ("📅 Thời điểm xuất file:", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
        ("📊 Nguồn dữ liệu:", "BCTN (14,528 file: Zenodo 13,982 + Supplement 546) | BCTC 702 chỉ tiêu chuẩn hóa | Báo chí kinh tế & Website DN"),
        ("📈 Cấu trúc dữ liệu:", "Chuẩn Dữ Liệu Bảng (Firm - Year Panel Data: Mã CK 'ticker', Năm 'year')"),
        ("🎓 Ứng dụng mô hình:", "Sẵn sàng hồi quy OLS, Fixed Effects, Random Effects, GMM (Stata, Python, R)"),
        ("🛡️ Mã xác thực dữ liệu:", "ARMINER-VERIFIED-PANEL-DATASET (Trương Minh Quân)"),
    ]

    for label, val in info_items:
        ws.row_dimensions[curr_row].height = 20
        c_lbl = ws.cell(row=curr_row, column=2, value=label)
        c_lbl.font = FONT_BODY_BOLD
        c_lbl.border = THIN_BORDER
        c_lbl.alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=4)
        c_val = ws.cell(row=curr_row, column=3, value=val)
        c_val.font = FONT_BODY
        for col in (3, 4):
            ws.cell(row=curr_row, column=col).border = THIN_BORDER
        c_val.alignment = Alignment(horizontal="left", vertical="center")
        curr_row += 1

    curr_row += 1

    # --- Section II: Mục lục các Sheet ---
    ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=4)
    sec2 = ws.cell(row=curr_row, column=2, value="II. MỤC LỤC & HƯỚNG DẪN SỬ DỤNG CÁC BẢNG DỮ LIỆU")
    sec2.font = Font(name="Segoe UI", size=11, bold=True, color=NAVY)
    sec2.fill = FILL_SECTION_HDR
    sec2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in (2, 3, 4):
        ws.cell(row=curr_row, column=col).border = THIN_BORDER
    ws.row_dimensions[curr_row].height = 24
    curr_row += 1

    other_sheets = [s for s in wb.sheetnames if s not in ("Trang_Bia", "Cover_Sheet", "Cover")]
    if not other_sheets:
        other_sheets = ["Firm_Year_Panel", "Articles_Detail", "Raw_Keywords"]

    for s_name in other_sheets:
        ws.row_dimensions[curr_row].height = 20
        c_name = ws.cell(row=curr_row, column=2, value=f"→ {s_name}")
        c_name.font = Font(name="Segoe UI", size=10, bold=True, color=NAVY)
        c_name.border = THIN_BORDER

        desc = _SHEET_DESCRIPTIONS.get(s_name, "Bảng dữ liệu trích xuất từ arminer Web Studio")
        ws.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=4)
        c_desc = ws.cell(row=curr_row, column=3, value=desc)
        c_desc.font = FONT_BODY
        for col in (3, 4):
            ws.cell(row=curr_row, column=col).border = THIN_BORDER
        curr_row += 1

    return ws


def style_worksheet(
    ws,
    header_fill: Optional[PatternFill] = None,
    freeze_at: Optional[str] = None,
    auto_filter: bool = True,
    auto_width: bool = True,
    tab_color: Optional[str] = None,
    number_formats: Optional[Dict[str, str]] = None,
) -> None:
    """
    Apply premium styling to a worksheet that already has data.

    Args:
        ws: openpyxl Worksheet with data written (row 1 = headers).
        header_fill: PatternFill for header row. Auto-picked if None.
        freeze_at: Cell ref for freeze panes (e.g. "C2"). Auto-set if None.
        auto_filter: Whether to add AutoFilter.
        auto_width: Whether to auto-size columns.
        tab_color: Sheet tab color hex. Auto-picked if None.
        number_formats: Dict of {column_name: format_string} overrides.
    """
    if ws.max_row is None or ws.max_row < 1:
        return

    # Skip Trang_Bia from raw data grid styling
    if ws.title in ("Trang_Bia", "Cover_Sheet", "Cover"):
        return

    max_col = ws.max_column or 1
    max_row = ws.max_row or 1

    # Tab color
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    else:
        color = _SHEET_COLORS.get(ws.title, NAVY)
        ws.sheet_properties.tabColor = color

    # Header fill
    if header_fill is None:
        header_fill = _pick_header_fill(ws.title)

    # Header row height
    ws.row_dimensions[1].height = 28

    # --- Style header row (row 1) ---
    col_names = []
    for col_idx in range(1, max_col + 1):
        cell = ws.cell(row=1, column=col_idx)
        cname = str(cell.value or "")
        col_names.append(cname)
        cell.font = FONT_HEADER
        cell.fill = header_fill
        cell.alignment = CENTER
        cell.border = HEADER_BORDER

    # --- Build number_formats map (col_index -> fmt) ---
    nf_map: Dict[int, str] = {}
    user_nf = number_formats or {}
    for col_idx, cname in enumerate(col_names, 1):
        if cname in user_nf:
            nf_map[col_idx] = user_nf[cname]
        else:
            guessed = _guess_number_format(cname)
            if guessed:
                nf_map[col_idx] = guessed

    # --- Style data rows (row 2+) ---
    for row_idx in range(2, max_row + 1):
        ws.row_dimensions[row_idx].height = 21

        # Zebra striping
        fill = FILL_WHITE if (row_idx - 2) % 2 == 0 else FILL_ZEBRA

        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = FONT_BODY
            cell.fill = fill
            cell.border = THIN_BORDER

            cname_lower = col_names[col_idx - 1].lower()

            # Alignments & Formats
            if cname_lower in ("ticker", "ma_ck", "code"):
                cell.font = FONT_BODY_BOLD
                cell.alignment = CENTER
            elif cname_lower in ("year", "nam", "published_year"):
                cell.alignment = CENTER
                cell.number_format = "0"
            elif cname_lower in ("mention", "dummy", "substantive"):
                cell.alignment = CENTER
                cell.number_format = "0"
            elif isinstance(cell.value, (int, float)):
                cell.alignment = RIGHT
                if col_idx in nf_map:
                    cell.number_format = nf_map[col_idx]
                else:
                    v = abs(cell.value)
                    if v == 0 or v >= 100:
                        cell.number_format = "#,##0"
                    elif v >= 1:
                        cell.number_format = "#,##0.00"
                    else:
                        cell.number_format = "0.0000"
            else:
                cell.alignment = LEFT
                if cname_lower == "url" and str(cell.value or "").startswith("http"):
                    cell.font = Font(name="Segoe UI", size=9.5, color="2563EB", underline="single")

    # --- Freeze panes ---
    if freeze_at:
        ws.freeze_panes = freeze_at
    elif ws.title == "Context":
        ws.freeze_panes = "D2"  # Freeze STT, Firm, Year
    elif ws.title in ("Descriptive_Stats", "Correlation"):
        ws.freeze_panes = "B2"  # Freeze Variable column
    elif max_col >= 3:
        ws.freeze_panes = "C2"  # Freeze Ticker, Year
    else:
        ws.freeze_panes = "A2"

    # --- AutoFilter ---
    if auto_filter and max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    # --- Page setup / footer ---
    try:
        ws.oddFooter.left.text = "arminer Studio | Trương Minh Quân"
        ws.oddFooter.center.text = "Hệ Thống Dữ Liệu Doanh Nghiệp & Kinh Tế Lượng"
        ws.oddFooter.right.text = "Trang &P / &N"
    except Exception:
        pass

    # --- Auto column width ---
    if auto_width:
        # Context sheet: wider max_width for Sentence_Context column
        if ws.title == "Context":
            _auto_col_width(ws, max_width=90)
            # Force Sentence_Context column to be extra wide with wrap_text
            for col_idx, cname in enumerate(col_names, 1):
                if cname.lower() in ("sentence_context",):
                    col_letter = get_column_letter(col_idx)
                    ws.column_dimensions[col_letter].width = 80
                    # Apply wrap_text to all data cells in this column
                    for row_idx in range(2, max_row + 1):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        cell.alignment = Alignment(
                            horizontal="left", vertical="center", wrap_text=True
                        )
                        ws.row_dimensions[row_idx].height = 36
        else:
            _auto_col_width(ws)


def style_workbook(wb: openpyxl.Workbook, custom_title: Optional[str] = None) -> openpyxl.Workbook:
    """Apply premium styling and embed Cover Sheet to ALL sheets in a workbook."""
    # 1. Add / Ensure Cover Sheet exists at index 0
    add_cover_sheet(wb, custom_title=custom_title)

    # 2. Style all data worksheets (even with 1 row header)
    for ws in wb.worksheets:
        if ws.title not in ("Trang_Bia", "Cover_Sheet", "Cover") and ws.max_row and ws.max_row >= 1:
            style_worksheet(ws)

    # 3. Set Document Core Properties (Brand Attribution)
    try:
        wb.properties.creator = "arminer Web Studio (Trương Minh Quân)"
        wb.properties.lastModifiedBy = "arminer Web Studio v2.5"
        wb.properties.title = custom_title or "arminer Research Dataset - Vietnam Listed Companies"
        wb.properties.subject = "Corporate Text Mining & Financial Econometric Panel Data"
        wb.properties.description = "Phát triển bởi Trương Minh Quân - DUE"
        wb.properties.category = "Econometric Panel Data"
    except Exception as e:
        logger.debug(f"Could not set workbook properties: {e}")

    return wb


def style_excel_file(path: Union[str, Path], custom_title: Optional[str] = None) -> Path:
    """
    Open an existing .xlsx file, apply premium styling & Cover Sheet, and save.
    Convenience function for post-processing after pd.ExcelWriter.

    Args:
        path: Path to the .xlsx file.
        custom_title: Optional custom title for cover sheet.

    Returns:
        The same Path after styling.
    """
    path = Path(path)
    if not path.exists():
        logger.warning(f"style_excel_file: file not found: {path}")
        return path

    try:
        wb = openpyxl.load_workbook(path)
        style_workbook(wb, custom_title=custom_title)
        wb.save(path)
        wb.close()
        logger.info(f"Applied premium styling & cover branding to {path.name}")
    except Exception as e:
        logger.warning(f"Could not apply styling to {path.name}: {e}")

    return path
